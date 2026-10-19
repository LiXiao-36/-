#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
龙洲校区「家作一周一上传」核心检查逻辑（CLI 与 GUI 共用）

不包含任何命令行或界面代码，只做：
  读配班表 → 枚举周目录 → 扫描文件判定班号 → 算未交 → 生成文案 / 导出表格
"""

from __future__ import annotations

import base64
import csv
import os
import re
import subprocess
import sys
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Callable, Iterable

# ============================ 可按需修改的默认配置 ============================

DEFAULT_YEAR_DIR = Path(r"Z:\6.教研室\☆龙洲校区 家作一周一上传\2026学年第一学期")
DEFAULT_ROSTER_XLS = Path(r"C:\Users\yungz\Desktop\2026学年龙洲教师配班安排表.xls")

BASE_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) \
    else Path(__file__).resolve().parent
ROSTER_CSV = BASE_DIR / "配班表.csv"
OUT_DIR = BASE_DIR / "输出"

# ==========================================================================

GRADE_NAMES = {
    "1": "一年级", "2": "二年级", "3": "三年级",
    "4": "四年级", "5": "五年级", "6": "六年级",
}
GRADE_ORDER = {v: int(k) for k, v in GRADE_NAMES.items()}

IGNORE_FILES = {"thumbs.db", "desktop.ini", ".ds_store"}
IGNORE_SUFFIX = {".db", ".tmp", ".lnk", ".ini"}

WEEK_DIR_RE = re.compile(r"第\s*(\d+)\s*周")
WEEK_RANGE_RE = re.compile(r"[（(]\s*([^）)]*?)\s*[）)]")


# ----------------------------- 小工具 -----------------------------

def setup_console_hint() -> None:
    """Windows 控制台尽量用 UTF-8 输出，避免中文乱码。"""
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
        sys.stderr.reconfigure(encoding="utf-8")  # type: ignore[attr-defined]
    except Exception:
        pass


def norm_name(s: str) -> str:
    """去掉所有空白（含全角空格）：'谢  晶' -> '谢晶'。"""
    return re.sub(r"\s+", "", s or "")


def cjk_width(s: str) -> int:
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in s)


def pad(s: str, width: int) -> str:
    return s + " " * max(0, width - cjk_width(s))


def fmt_time(dt: datetime) -> str:
    return f"{dt.month}月{dt.day}日 {dt:%H:%M}"


def week_label(week_no: int | None, week_dir: Path) -> str:
    m = WEEK_RANGE_RE.search(week_dir.name)
    rng = m.group(1) if m else ""
    base = f"第{week_no}周" if week_no is not None else week_dir.name
    return f"{base}（{rng}）" if rng else base


def share_path_label(year_dir: Path) -> str:
    """用于文案里的共享盘路径，如『☆龙洲校区 家作一周一上传\\2026学年第一学期』。"""
    parent = year_dir.parent.name
    if not parent or parent.endswith(":"):
        return year_dir.name
    return f"{parent}\\{year_dir.name}"


# ----------------------------- 配班表 -----------------------------

PS_TEMPLATE = r"""
$ErrorActionPreference = 'Stop'
$xls = '@@XLS@@'
$out = '@@OUT@@'
$xl = New-Object -ComObject Excel.Application
$xl.Visible = $false
$xl.DisplayAlerts = $false
$wb = $null
try {
    $wb = $xl.Workbooks.Open($xls, 0, $true)
    $ws = $wb.Sheets.Item(1)
    $rng = $ws.UsedRange
    $rows = $rng.Rows.Count
    $cols = $rng.Columns.Count
    $vals = $rng.Value2

    function Get-Cell([object]$grid, [int]$r, [int]$c) {
        $v = $grid.GetValue($r, $c)
        if ($null -eq $v) { return '' }
        return [string]$v
    }

    $headerRow = 0
    $classCols = @()
    for ($r = 1; $r -le $rows; $r++) {
        $hit = @()
        for ($c = 1; $c -le $cols; $c++) {
            if ((Get-Cell $vals $r $c).Trim() -eq '班级') { $hit += $c }
        }
        if ($hit.Count -gt 0) { $headerRow = $r; $classCols = $hit; break }
    }
    if ($headerRow -eq 0) { throw '配班表中找不到表头「班级」' }

    $lines = [System.Collections.Generic.List[string]]::new()
    $lines.Add('班级,班主任,配班老师')
    $seen = @{}
    for ($r = $headerRow + 1; $r -le $rows; $r++) {
        foreach ($c in $classCols) {
            $code = (Get-Cell $vals $r $c) -replace '\s', ''
            if ($code -notmatch '^\d{3}$') { continue }
            if ($seen.ContainsKey($code)) { continue }
            $seen[$code] = $true
            $t = ((Get-Cell $vals $r ($c + 1)) -replace '\s', '')
            $p = ((Get-Cell $vals $r ($c + 2)) -replace '\s', '')
            $lines.Add("$code,$t,$p")
        }
    }
    [IO.File]::WriteAllLines($out, $lines, [System.Text.UTF8Encoding]::new($true))
    Write-Output ("OK " + $lines.Count)
}
finally {
    if ($null -ne $wb) { $wb.Close($false) }
    $xl.Quit()
    [void][Runtime.InteropServices.Marshal]::ReleaseComObject($xl)
    [GC]::Collect()
}
"""


def _powershell_exe() -> str | None:
    for name in ("pwsh", "powershell"):
        try:
            subprocess.run([name, "-NoProfile", "-Command", "$PSVersionTable.PSVersion.Major"],
                           capture_output=True, timeout=30, check=True)
            return name
        except Exception:
            continue
    return None


def convert_roster_xls(xls: Path, out_csv: Path) -> None:
    """用 Excel COM 把 .xls 配班表转成 UTF-8 CSV（csv 列：班级,班主任,配班老师）。"""
    exe = _powershell_exe()
    if exe is None:
        raise RuntimeError("找不到可用的 PowerShell")
    code = PS_TEMPLATE.replace("@@XLS@@", str(xls)).replace("@@OUT@@", str(out_csv))
    b64 = base64.b64encode(code.encode("utf-16-le")).decode("ascii")
    proc = subprocess.run(
        [exe, "-NoProfile", "-NonInteractive", "-EncodedCommand", b64],
        capture_output=True, timeout=180,
    )
    if proc.returncode != 0 or not out_csv.exists():
        err = (proc.stderr or b"").decode("utf-8", "ignore").strip()
        out = (proc.stdout or b"").decode("utf-8", "ignore").strip()
        raise RuntimeError(f"读取 .xls 配班表失败：{err or out or '未知错误'}")


def _grid_to_rows(grid: list[list[str]]) -> list[tuple[str, str, str]]:
    """横向双栏表 → [(班号, 班主任, 配班老师)]。"""
    header_row = None
    class_cols: list[int] = []
    for r, row in enumerate(grid):
        hits = [c for c, v in enumerate(row) if (v or "").strip() == "班级"]
        if hits:
            header_row, class_cols = r, hits
            break
    if header_row is None:
        raise RuntimeError("配班表中找不到表头「班级」")

    out: list[tuple[str, str, str]] = []
    seen: set[str] = set()
    for row in grid[header_row + 1:]:
        for c in class_cols:
            raw = (row[c] if c < len(row) else "") or ""
            code = norm_name(raw)
            if not re.fullmatch(r"\d{3}", code) or code in seen:
                continue
            seen.add(code)

            def cell(i: int) -> str:
                j = c + i
                return norm_name(row[j]) if j < len(row) else ""

            out.append((code, cell(1), cell(2)))
    return out


def convert_roster_xlsx(xlsx: Path, out_csv: Path) -> None:
    """纯 Python 读 .xlsx（不依赖 Excel）。"""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    grid: list[list[str]] = []
    for row in ws.iter_rows(values_only=True):
        grid.append(["" if v is None else str(v) for v in row])
    wb.close()

    rows = _grid_to_rows(grid)
    if not rows:
        raise RuntimeError("配班表里没有读到任何班级")
    with open(out_csv, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["班级", "班主任", "配班老师"])
        w.writerows(rows)


def load_roster(roster_path: Path, csv_path: Path = ROSTER_CSV,
                force: bool = False, log: Callable[[str], None] | None = None) -> dict[str, dict]:
    """返回 {班号: {grade, teacher, partner}}。.xls 走 Excel COM，.xlsx 走 openpyxl。"""
    say = log or (lambda _m: None)
    need = force or not csv_path.exists()
    if not need and roster_path.exists():
        try:
            need = roster_path.stat().st_mtime > csv_path.stat().st_mtime
        except OSError:
            need = False

    if need:
        if roster_path.exists():
            try:
                suf = roster_path.suffix.lower()
                if suf == ".xlsx":
                    convert_roster_xlsx(roster_path, csv_path)
                else:
                    convert_roster_xls(roster_path, csv_path)
                say(f"配班表已从 {roster_path.name} 刷新")
            except Exception as e:
                if csv_path.exists():
                    say(f"读取配班表失败（{e}），改用缓存 {csv_path.name}")
                else:
                    raise
        elif csv_path.exists():
            say(f"找不到 {roster_path.name}，改用缓存 {csv_path.name}")
        else:
            raise FileNotFoundError(f"配班表不存在：{roster_path}")

    roster: dict[str, dict] = {}
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            code = norm_name(row.get("班级", ""))
            if not re.fullmatch(r"\d{3}", code):
                continue
            roster[code] = {
                "grade": (row.get("年级") or GRADE_NAMES.get(code[0], "")).strip()
                         or GRADE_NAMES.get(code[0], ""),
                "teacher": norm_name(row.get("班主任", "")),
                "partner": norm_name(row.get("配班老师", "")),
            }
    if not roster:
        raise RuntimeError("配班表里没有读到任何班级，请检查表格格式")
    return roster


# ----------------------------- 周目录 -----------------------------

def list_weeks(year_dir: Path) -> list[tuple[int, Path]]:
    """返回 [(周次, 路径)]，按周次升序。若本目录本身就是某一周，返回单元素。"""
    if not year_dir.exists():
        raise FileNotFoundError(f"目录不存在或无法访问：{year_dir}")

    self_match = WEEK_DIR_RE.search(year_dir.name)
    if self_match:
        return [(int(self_match.group(1)), year_dir)]

    items: list[tuple[int, Path]] = []
    for p in year_dir.iterdir():
        if not p.is_dir():
            continue
        m = WEEK_DIR_RE.search(p.name)
        if m:
            items.append((int(m.group(1)), p))
    items.sort(key=lambda x: x[0])
    return items


def count_files(week_dir: Path) -> int:
    return sum(1 for _r, _d, fs in os.walk(week_dir) for f in fs if not should_skip(f))


# ----------------------------- 扫描判定 -----------------------------

def should_skip(name: str) -> bool:
    low = name.lower()
    if low in IGNORE_FILES or low.startswith("~$"):
        return True
    return Path(low).suffix in IGNORE_SUFFIX


def extract_code(rel_parts: tuple[str, ...], filename: str, valid: set[str]) -> str | None:
    """
    推断文件属于哪个班，优先级：
      1) 子文件夹名正好是班号
      2) 文件名开头三位数（301.JPG / 204班1--2.jpg / 402-1.jpg）
      3) 文件名里任意三位数且该班号在配班表中
    """
    for part in rel_parts[:-1]:
        if part in valid:
            return part

    stem = Path(filename).stem
    m = re.match(r"^(\d{3})(?!\d)", stem)
    if m and m.group(1) in valid:
        return m.group(1)

    for tok in re.findall(r"(?<!\d)(\d{3})(?!\d)", stem):
        if tok in valid:
            return tok
    return None


def scan_submissions(week_dir: Path, valid: set[str]) -> tuple[dict[str, list[str]], list[str]]:
    submitted: dict[str, list[str]] = {}
    unknown: list[str] = []
    for root, dirs, files in os.walk(week_dir):
        dirs[:] = [d for d in dirs if not d.startswith(".")]
        for fn in files:
            if should_skip(fn):
                continue
            full = Path(root) / fn
            try:
                rel = full.relative_to(week_dir)
            except ValueError:
                rel = Path(fn)
            code = extract_code(rel.parts, fn, valid)
            if code:
                submitted.setdefault(code, []).append(str(rel))
            else:
                unknown.append(str(rel))
    return submitted, unknown


# ----------------------------- 单周结果 -----------------------------

@dataclass
class WeekResult:
    week_no: int | None
    path: Path
    label: str
    scanned_at: datetime
    total: int = 0
    submitted: dict[str, list[str]] = field(default_factory=dict)
    unknown: list[str] = field(default_factory=list)
    missing: list[dict] = field(default_factory=list)
    submitted_rows: list[dict] = field(default_factory=list)
    zero_grades: list[str] = field(default_factory=list)

    @property
    def ok_count(self) -> int:
        return len(self.submitted_rows)

    @property
    def missing_count(self) -> int:
        return len(self.missing)

    def status_of(self, code: str) -> str:
        return "ok" if self.submitted.get(code) else "no"


def analyze_week(roster: dict[str, dict], week_no: int | None, week_dir: Path) -> WeekResult:
    submitted, unknown = scan_submissions(week_dir, set(roster))
    res = WeekResult(
        week_no=week_no,
        path=week_dir,
        label=week_label(week_no, week_dir),
        scanned_at=datetime.now(),
        total=len(roster),
        submitted=submitted,
        unknown=unknown,
    )

    for code in sorted(roster):
        info = roster[code]
        files = submitted.get(code, [])
        row = {
            "code": code,
            "grade": info["grade"],
            "teacher": info["teacher"],
            "partner": info["partner"],
            "files": files,
        }
        if files:
            res.submitted_rows.append(row)
        else:
            row["note"] = "配班表无班主任姓名" if not info["teacher"] else ""
            res.missing.append(row)

    order = lambda r: (GRADE_ORDER.get(r["grade"], 99), r["code"])
    res.missing.sort(key=order)
    res.submitted_rows.sort(key=order)

    for g in GRADE_ORDER:
        total_g = sum(1 for r in roster.values() if r["grade"] == g)
        ok_g = sum(1 for r in res.submitted_rows if r["grade"] == g)
        if total_g >= 3 and ok_g == 0:
            res.zero_grades.append(g)

    return res


def analyze_weeks(roster: dict[str, dict],
                  weeks: Iterable[tuple[int | None, Path]],
                  progress: Callable[[int, int, str], None] | None = None,
                  should_stop: Callable[[], bool] | None = None) -> list[WeekResult]:
    pairs = list(weeks)
    out: list[WeekResult] = []
    for i, (no, path) in enumerate(pairs, 1):
        if should_stop and should_stop():
            break
        if progress:
            progress(i, len(pairs), path.name)
        out.append(analyze_week(roster, no, path))
    return out


# ----------------------------- 文案 -----------------------------

def build_single_message(code: str, info: dict, label: str, when: str,
                         share: str) -> str:
    teacher = info["teacher"] or "老师"
    partner = info.get("partner") or ""
    tail = f"（配班：{partner}）" if partner else ""
    return (
        f"{teacher}老师，您好！\n"
        f"截至{when}，{code}班{label}的家作还未上传到共享盘“{share}\\{label}”。\n"
        f"{tail}麻烦您抽空补交一下，谢谢！"
    )


def build_group_message(missing: list[dict], label: str, when: str, share: str) -> str:
    lines = [
        f"【{label}家作上传提醒】",
        f"截至{when}，以下 {len(missing)} 个班级尚未上传家作，请对应班主任今天内补交：",
        "",
    ]
    by_grade: dict[str, list[dict]] = {}
    for row in missing:
        by_grade.setdefault(row["grade"], []).append(row)
    for g, rows in by_grade.items():
        names = "、".join(f"{r['code']}（{r['teacher']}）" for r in rows)
        lines.append(f"{g} {len(rows)}个班：{names}")
    lines += ["", f"上传路径：\\\\10.178.40.10\\完全共享\\6.教研室\\{share}"]
    return "\n".join(lines)


def messages_for(result: WeekResult, mode: str, year_dir: Path,
                 when: str | None = None) -> str:
    """mode: 'group' 群发一条 / 'single' 逐个发送。"""
    share = share_path_label(year_dir)
    when = when or fmt_time(result.scanned_at)
    if mode == "group":
        return build_group_message(result.missing, result.label, when, share)
    blocks = []
    for row in result.missing:
        blocks.append("═" * 46)
        blocks.append(f"【{row['code']}班  {row['teacher'] or '（配班表无姓名）'}】")
        blocks.append("─" * 46)
        blocks.append(build_single_message(row["code"], row, result.label, when, share))
        blocks.append("")
    return "\n".join(blocks)


# ----------------------------- 导出 -----------------------------

def write_csv(path: Path, rows: list[dict]) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["序号", "年级", "班级", "班主任", "配班老师", "已交文件数", "备注"])
        for i, r in enumerate(rows, 1):
            w.writerow([i, r["grade"], r["code"], r["teacher"], r["partner"],
                        len(r.get("files") or []), r.get("note", "")])


def write_messages(out_dir: Path, tag: str, result: WeekResult, year_dir: Path) -> tuple[Path, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    single = out_dir / f"钉钉_逐个发送_{tag}.txt"
    group = out_dir / f"钉钉_群发一条_{tag}.txt"
    single.write_text(messages_for(result, "single", year_dir), encoding="utf-8-sig")
    group.write_text(messages_for(result, "group", year_dir), encoding="utf-8-sig")
    return single, group


def write_xlsx(path: Path, result: WeekResult, year_dir: Path) -> None:
    """单周明细表：未交清单 / 已交清单 / 未识别文件。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "未交清单"

    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws.merge_cells("A1:G1")
    ws["A1"] = f"{result.label} 家作未交清单"
    ws["A1"].font = Font(bold=True, size=15, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:G2")
    ws["A2"] = (f"扫描目录：{result.path}\n"
                f"统计时间：{result.scanned_at:%Y-%m-%d %H:%M}    "
                f"班级总数：{result.total}    已交：{result.ok_count}    未交：{result.missing_count}")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 34

    headers = ["序号", "年级", "班级", "班主任", "配班老师", "已交文件数", "备注"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=i, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[3].height = 24

    r = 4
    for i, row in enumerate(result.missing, 1):
        vals = [i, row["grade"], row["code"], row["teacher"], row["partner"],
                len(row.get("files") or []), row.get("note", "")]
        for j, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=j, value=v)
            cell.border = border
            cell.alignment = left if j == 7 else center
            if j == 4:
                cell.font = Font(bold=True, color="C00000")
            if r % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
        r += 1

    for i, w in enumerate([6, 10, 9, 12, 12, 12, 26], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    if result.missing:
        ws.auto_filter.ref = f"A3:G{r - 1}"

    ws2 = wb.create_sheet("已交清单")
    for i, h in enumerate(["序号", "年级", "班级", "班主任", "配班老师", "文件"], 1):
        cell = ws2.cell(row=1, column=i, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = center
        cell.border = border
    ws2.row_dimensions[1].height = 24
    r2 = 2
    for i, row in enumerate(result.submitted_rows, 1):
        vals = [i, row["grade"], row["code"], row["teacher"], row["partner"],
                "\n".join(row.get("files") or [])]
        for j, v in enumerate(vals, 1):
            cell = ws2.cell(row=r2, column=j, value=v)
            cell.border = border
            cell.alignment = left if j == 6 else center
        r2 += 1
    for i, w in enumerate([6, 10, 9, 12, 12, 60], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    if result.unknown:
        ws3 = wb.create_sheet("未识别文件")
        ws3["A1"] = "以下文件无法从名称判断班号，请人工核对："
        ws3["A1"].font = Font(bold=True, color="C00000")
        for i, u in enumerate(result.unknown, 1):
            ws3.cell(row=i + 2, column=1, value=u)
        ws3.column_dimensions["A"].width = 90

    ws.sheet_view.showGridLines = False
    wb.save(path)


def write_matrix_xlsx(path: Path, results: list[WeekResult], roster: dict[str, dict],
                      year_dir: Path, codes: list[str] | None = None,
                      note: str = "") -> None:
    """多周总表：总览矩阵 + 未交明细。

    codes 不为空时，只导出这些班级（界面「导出筛选结果」用）。
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    if codes:
        order = [c for c in codes if c in roster]
    else:
        order = sorted(roster, key=lambda c: (GRADE_ORDER.get(roster[c]["grade"], 99), c))
    only = set(order)

    wb = Workbook()
    ws = wb.active
    ws.title = "总览矩阵" if not codes else "筛选结果"

    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    ok_fill = PatternFill("solid", fgColor="E2EFDA")
    no_fill = PatternFill("solid", fgColor="FCE4E4")
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    ncols = 3 + len(results) + 1
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title = f"家作上交总览　共 {len(results)} 周　{year_dir.name}"
    if codes:
        title += f"　（筛选 {len(order)} 个班）"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=15, color="1F4E79")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    sub = f"导出时间：{datetime.now():%Y-%m-%d %H:%M}　　√=已交　×=未交"
    if note:
        sub += f"　　{note}"
    ws["A2"] = sub
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[2].height = 20

    headers = ["班级", "班主任", "配班老师"] + [r.label for r in results] + ["未交次数"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=i, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.row_dimensions[3].height = 42

    r = 4
    for code in order:
        info = roster[code]
        ws.cell(row=r, column=1, value=code).border = border
        ws.cell(row=r, column=2, value=info["teacher"]).border = border
        ws.cell(row=r, column=3, value=info["partner"]).border = border
        miss = 0
        for j, res in enumerate(results, 4):
            ok = bool(res.submitted.get(code))
            cell = ws.cell(row=r, column=j, value="√" if ok else "×")
            cell.alignment = center
            cell.border = border
            cell.fill = ok_fill if ok else no_fill
            cell.font = Font(color="2E7D32" if ok else "C00000", bold=True)
            if not ok:
                miss += 1
        cell = ws.cell(row=r, column=4 + len(results), value=miss)
        cell.alignment = center
        cell.border = border
        if miss:
            cell.font = Font(bold=True, color="C00000")
        r += 1

    for i in range(1, 4):
        for rr in range(4, r):
            ws.cell(row=rr, column=i).alignment = center
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    for i in range(4, 4 + len(results)):
        ws.column_dimensions[get_column_letter(i)].width = 13
    ws.column_dimensions[get_column_letter(4 + len(results))].width = 10
    ws.freeze_panes = "D4"

    # 未交明细
    ws2 = wb.create_sheet("未交明细")
    for i, h in enumerate(["周次", "年级", "班级", "班主任", "配班老师"], 1):
        cell = ws2.cell(row=1, column=i, value=h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = center
        cell.border = border
    rr = 2
    for res in results:
        for row in res.missing:
            if codes and row["code"] not in only:
                continue
            vals = [res.label, row["grade"], row["code"], row["teacher"], row["partner"]]
            for j, v in enumerate(vals, 1):
                cell = ws2.cell(row=rr, column=j, value=v)
                cell.border = border
                cell.alignment = center
            rr += 1
    for i, w in enumerate([20, 10, 9, 12, 12], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"

    ws.sheet_view.showGridLines = False
    wb.save(path)
